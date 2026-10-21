from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from ..rt import OUTSIDE
from ..spec import TaskSpec
from ..text import markup_tags

TABLE_SUFFIXES = (".csv", ".xlsx")
TEMPLATE_STEM = "template"
SPLIT_CHOICES = ["train", "test"]
LABEL_COL, SPLIT_COL = "label", "split"
DATA_SHEET, GUIDE_SHEET = "data", "labels"
MAX_DROPDOWN_ROWS = 100000


def is_table(path: str | Path | None) -> bool:
    return bool(path) and Path(path).suffix.lower() in TABLE_SUFFIXES


def input_column(spec: TaskSpec) -> str:
    return "text" if spec.modality == "text" else "file"


def columns(spec: TaskSpec) -> list[str]:
    if spec.task_type == "frame" and spec.modality != "text":
        raise ValueError("per-frame audio labels do not fit a table: give timed segment files instead")
    tagged = spec.task_type == "frame"
    return [input_column(spec), *([] if tagged else [LABEL_COL]), SPLIT_COL]


def read_table(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [{k.strip(): (v or "").strip() for k, v in r.items() if k} for r in csv.DictReader(f)]
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)[DATA_SHEET]
    rows = ws.iter_rows(values_only=True)
    head = [str(h or "").strip() for h in next(rows, [])]
    return [{h: str(v or "").strip() for h, v in zip(head, r, strict=False) if h} for r in rows if any(r)]


def _write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def _dropdown(ws: Any, col: int, choices: list[str], sheet: str) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ref = f"'{sheet}'!$A$2:$A${len(choices) + 1}" if sheet else '"' + ",".join(choices) + '"'
    dv = DataValidation(type="list", formula1=ref, allow_blank=True)
    letter = get_column_letter(col)
    dv.add(f"{letter}2:{letter}{MAX_DROPDOWN_ROWS}")
    ws.add_data_validation(dv)


def _write_xlsx(path: Path, header: list[str], rows: list[list[Any]], guide: list[list[Any]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    ws.append(header)
    for r in rows:
        ws.append(r)
    labels = wb.create_sheet(GUIDE_SHEET)
    for r in guide:
        labels.append(r)
    if LABEL_COL in header:
        _dropdown(ws, header.index(LABEL_COL) + 1, [r[0] for r in guide[1:]], GUIDE_SHEET)
    _dropdown(ws, header.index(SPLIT_COL) + 1, SPLIT_CHOICES, "")
    wb.save(path)


def write_template(out: Path, header: list[str], rows: list[list[Any]], guide: list[list[Any]]) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    csv_path, xlsx_path = (out / f"{TEMPLATE_STEM}{s}" for s in TABLE_SUFFIXES)
    _write_csv(csv_path, header, rows)
    _write_xlsx(xlsx_path, header, rows, guide)
    return [csv_path, xlsx_path]


def resolve_input(spec: TaskSpec, raw: str, base: Path) -> str:
    if spec.modality == "text":
        return raw
    p = Path(raw).expanduser()
    return str((p if p.is_absolute() else base / p).resolve())


def table_rows(spec: TaskSpec, path: Path) -> list[tuple[Any, Any, str]]:
    out = []
    col = input_column(spec)
    for r in read_table(path):
        raw = r.get(col, "")
        if not raw:
            continue
        if spec.task_type == "frame":
            words, tags = markup_tags(raw)
            known = set(spec.label_names)
            out.append((" ".join(words), [t if t in known else OUTSIDE for t in tags], r.get(SPLIT_COL, "")))
        elif r.get(LABEL_COL) in spec.label_names:
            out.append((resolve_input(spec, raw, path.parent), r[LABEL_COL], r.get(SPLIT_COL, "")))
    return out
